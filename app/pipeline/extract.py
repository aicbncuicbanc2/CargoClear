"""Stage 2: extract the 7 shipment fields from SI / BL attachment text.

Attachments are a mix of formats (~78% .txt, ~22% real binary: pdf+pdf,
xlsx+docx, xlsx+xlsx — see data/attachments). `read_document()` normalizes
every format down to label/value text lines, so exactly one alias matcher
serves all of them.

Label synonymy is deliberate in this dataset (SI vs BL use different labels
for the same field) — confirmed by manually tracing email_004:

    field            | SI label                      | BL label
    -----------------|--------------------------------|---------------------------
    shipper          | Shipper                        | SHIPPER
    consignee        | Consignee (Non-Negotiable)      | To the Order of
    notify_party      | Notify                         | Notify Party
    port_of_loading   | Port of Loading (POL)          | Port of Loading (POL)
    port_of_discharge | POD                            | POD
    container_count   | Total Containers               | Container Count
    gross_weight_kg   | Gross Wt (kgs)                 | Gross Weight (KG)

Also per the dataset README, expect: "Load Port" as another POL alias.

Matching is by anchored label *prefix*: an alias must start the line and be
followed by a separator, then the rest of the line is the value. Anchoring
is what keeps short aliases like "pol"/"pod" from firing inside unrelated
labels, and the prefix form is required because the real documents are not
uniformly colon-delimited — the .txt/.xlsx pairs write
"Load Port: SINGAPORE" while the PDF pairs write "Load Port BUATAN,
INDONESIA" with a single space. Longest alias wins, so
"Consignee (Non-Negotiable)" beats a bare "Consignee".

Values are taken as the first line only. Shipper/consignee blocks in the
PDF pairs run on for several address lines, and SI and BL wrap those
addresses differently for the same party — comparing whole blocks would
manufacture mismatches, while the identity line that actually carries the
defect is always first.
"""

from __future__ import annotations

import re
from pathlib import Path

from app.pipeline.models import COMPARISON_FIELDS


class UnreadableDocument(Exception):
    """Raised when a file yields no usable text (scanned image, 0 bytes,
    garbled). Stage 4 turns this into review_reason='unreadable'."""


# Known label synonyms seen across SI/BL documents, listed canonical-first
# within each field (see _ALIAS_INDEX on why order matters).
#
# Every alias marked "observed" was harvested from the real attachments in
# data/attachments by scanning all 242 readable documents for labels the
# alias set did not yet cover. The unmarked ones are defensive variants.
#
# Deliberately NOT aliased: "Place of Receipt" / "Place of Delivery" (inland
# points, not the load/discharge ports — semantically different fields) and
# "Net Weight" (not gross). None of the three occurs in this dataset, so
# including them would add false-match risk with no recall to gain.
FIELD_ALIASES: dict[str, list[str]] = {
    "shipper": [
        "shipper (principal or seller)",  # observed: .txt, .xlsx
        "shipper",  # observed: .txt, .pdf
        "shippers",
        "shipper name",
        "shipper exporter",
        "exporter",
        "consignor",
    ],
    "consignee": [
        "consignee (non-negotiable)",  # observed: .txt, .pdf
        "to the order of",  # observed: .txt
        "consignee",  # observed: .txt, .xlsx
        "to order of",
        "to order",
        "order of",
        "consigned to",
        "consignee name",
    ],
    "notify_party": [
        "notify party/intermediate consignee",  # observed: .txt
        "notify party",  # observed: .txt, .pdf, .xlsx
        "notify",  # observed: .txt
        "notify parties",
        "notify address",
        "party to notify",
        "also notify",
    ],
    "port_of_loading": [
        "port of loading (pol)",  # observed: .txt
        "port of loading",  # observed: .txt
        "load port",  # observed: .txt, .pdf, .xlsx
        "pol",
        "loading port",
        "port of load",
    ],
    "port_of_discharge": [
        "port of discharge (pod)",  # observed: .txt
        "port of discharge",  # observed: .txt, .pdf
        "discharge port",
        "pod",  # observed: .txt
        "port of unloading",
        "destination port",
    ],
    "container_count": [
        "no. of containers or packages",  # observed: .txt, .xlsx
        "total containers",  # observed: .txt
        "container count",  # observed: .txt, .pdf
        "number of containers",
        "no of containers",
        "qty of containers",
        "container qty",
        "total container",
        "containers",
    ],
    "gross_weight_kg": [
        # Observed verbatim in 8 of the 20 readable PDFs as
        # "TOTAL Gross Weightnn(KGS):" — the generator literalized an "\n\n"
        # escape into the PDF text layer. Matched as written rather than by
        # loosening the matcher, which would weaken every other alias.
        "total gross weightnn (kgs)",
        "total gross weight (kg)",  # observed: .pdf
        "total gross wt (kgs)",  # observed: .pdf
        "total gross weight",  # observed: .pdf
        "gross weight (kg)",  # observed: .txt, .xlsx
        "gross wt (kgs)",  # observed: .txt
        "gross weight",
        "gross wt",
        "gross weight kgs",
        "gross mass",
    ],
}

assert set(FIELD_ALIASES) == set(COMPARISON_FIELDS)

# Values that mean "this was left blank", not "this is the value". These
# drive review_reason='missing_value' in stage 4 — never a mismatch.
BLANK_MARKERS = {
    "",
    "-",
    "--",
    "---",
    "n/a",
    "na",
    "nil",
    "none",
    "tba",
    "tbd",
    "tbc",
    "to be advised",
    "to be confirmed",
    "pending",
    "unknown",
    "?",
    "??",
    "???",
    "xxx",
    "xxxx",
}

# Separators allowed between a label and its value. An apostrophe is
# deliberately absent so "Shipper's Reference" is not read as "Shipper".
_SEPARATOR_CLASS = r"[\s:;,\-–—()\[\]]"
# Everything between the end of a label and the start of its value, consumed
# repeatedly in either form:
#   - a whole parenthetical, which is how the xlsx+docx pairs carry their
#     bilingual gloss: "Shipper (Principal or Seller) (发货人):"
#   - a run of separator punctuation, including the closing bracket left over
#     when an alias matched through "(POL" in "Port of Loading (POL):"
# An opening bracket is deliberately excluded from the separator run so that
# a parenthetical is always consumed whole by the first branch, never split.
_CONSUMED_SEPARATORS = r"(?:\s*\([^)]*\)|[\s:;,.\-–—)\]]+)*"


def normalize_label(label: str) -> str:
    """Flatten a label to its comparable form: lowercase, punctuation to
    spaces, whitespace collapsed. 'Gross Wt (kgs)' -> 'gross wt kgs'."""
    flattened = re.sub(r"[^0-9a-z]+", " ", label.lower())
    return re.sub(r"\s+", " ", flattened).strip()


# normalized alias -> (field, rank). Rank is the alias's position in its
# list: lists are canonical-first, so when one document uses two aliases of
# the same field the better-ranked label wins regardless of page order.
_ALIAS_INDEX: dict[str, tuple[str, int]] = {}
for _field, _aliases in FIELD_ALIASES.items():
    for _rank, _alias in enumerate(_aliases):
        _ALIAS_INDEX[normalize_label(_alias)] = (_field, _rank)


def _alias_pattern(normalized_alias: str) -> re.Pattern[str]:
    """Anchored prefix matcher for one alias.

    The alias's words are rejoined with a flexible punctuation separator so
    one entry matches "Port of Loading (POL)", "PORT OF LOADING - POL" and
    "port_of_loading" alike.
    """
    words = r"[^0-9A-Za-z]+".join(re.escape(word) for word in normalized_alias.split())
    return re.compile(
        rf"^\s*{words}(?={_SEPARATOR_CLASS}|$){_CONSUMED_SEPARATORS}(?P<value>.*)$",
        re.IGNORECASE,
    )


# Longest alias first so "Consignee (Non-Negotiable)" wins over "Consignee";
# ties broken by rank. Built once at import.
_ALIAS_PATTERNS: list[tuple[re.Pattern[str], str, int]] = [
    (_alias_pattern(alias), field, rank)
    for alias, (field, rank) in sorted(
        _ALIAS_INDEX.items(), key=lambda item: (-len(item[0]), item[1][1])
    )
]


def match_label(line: str) -> tuple[str, int, str] | None:
    """Match one line against the alias table.

    Returns (field, alias_rank, value) or None. Value may be empty when the
    label sits alone on its line.
    """
    for pattern, field, rank in _ALIAS_PATTERNS:
        found = pattern.match(line)
        if found:
            return field, rank, found.group("value").strip()
    return None


def is_blank_value(value: str | None) -> bool:
    """True when a value is present in the document but carries no content
    (blank line, '???', '______', 'TBA')."""
    if value is None:
        return True
    stripped = value.strip()
    if stripped.lower() in BLANK_MARKERS:
        return True
    # Runs of fill characters used as write-in blanks: ____, ....., ----
    return bool(re.fullmatch(r"[_.\-*·\s]+", stripped))


def is_present(fields: dict[str, str], name: str) -> bool:
    """The document mentions this field at all (blank or not)."""
    return name in fields


def is_blank(fields: dict[str, str], name: str) -> bool:
    """The document has the label but left the value blank — the condition
    behind review_reason='missing_value'."""
    return fields.get(name, None) == ""


def is_usable(fields: dict[str, str], name: str) -> bool:
    """There is a real value here to compare against."""
    return bool(fields.get(name))


def blank_fields(fields: dict[str, str]) -> list[str]:
    """Comparison fields present in the document but left blank."""
    return [name for name in COMPARISON_FIELDS if is_blank(fields, name)]


def _looks_like_label(line: str) -> bool:
    """Does this line start a new labelled field (rather than continue the
    previous value)? Used to stop multi-line value capture."""
    return match_label(line) is not None


def extract_fields(document_text: str) -> dict[str, str]:
    """Pull the 7 comparison fields out of one document's text.

    The return value is tri-state, because stage 4 has to tell "the document
    never mentions this field" apart from "the document has the field and
    left it blank" — only the latter is review_reason='missing_value':

        key absent          the label does not appear in the document
        key present, ""     the label appears but the value is blank
                            (a write-in blank, "???", "TBA", "N/A")
        key present, text   the extracted value

    Use `is_present`/`is_blank`/`is_usable` rather than truth-testing the
    value, so an empty string is never mistaken for an absent field.

    The first occurrence of a field wins — shipping documents repeat labels
    in footers and continuation pages.
    """
    fields: dict[str, str] = {}
    best_rank: dict[str, int] = {}

    lines = document_text.splitlines()
    for index, line in enumerate(lines):
        matched = match_label(line)
        if matched is None:
            continue
        field, rank, value = matched

        # First occurrence wins (documents repeat labels in footers and
        # continuation pages) unless this label is a better-ranked alias.
        if field in best_rank and rank >= best_rank[field]:
            continue

        # Block layout: "Consignee:" alone on its line, value beneath it.
        # Only the first continuation line is taken, matching the
        # first-line rule used when the value is inline.
        if not value:
            following = lines[index + 1].strip() if index + 1 < len(lines) else ""
            if following and not _looks_like_label(following):
                value = following

        best_rank[field] = rank
        # "" records "label found, value blank" — distinct from the key
        # being absent entirely. See this function's docstring.
        fields[field] = "" if is_blank_value(value) else value

    return fields


# --- format readers -------------------------------------------------------
# Each reader returns label/value text lines so extract_fields() can stay
# format-agnostic.


def read_document(path: str | Path) -> str:
    """Read any supported attachment into text.

    Raises UnreadableDocument for empty files, image-only scanned PDFs, and
    anything that yields no extractable text — stage 4 maps that to
    review_reason='unreadable'.
    """
    path = Path(path)
    if not path.exists():
        raise UnreadableDocument(f"{path} does not exist")
    if path.stat().st_size == 0:
        raise UnreadableDocument(f"{path} is empty (0 bytes)")

    suffix = path.suffix.lower()
    if suffix in {".txt", ".text", ".csv", ".md"}:
        text = path.read_text(encoding="utf-8", errors="replace")
    elif suffix == ".pdf":
        text = _read_pdf(path)
    elif suffix in {".docx", ".doc"}:
        text = _read_docx(path)
    elif suffix in {".xlsx", ".xlsm", ".xls"}:
        text = _read_xlsx(path)
    else:
        raise UnreadableDocument(f"{path}: unsupported format '{suffix}'")

    if not text or not text.strip():
        raise UnreadableDocument(f"{path}: no extractable text (scanned or garbled?)")
    return text


def _read_pdf(path: Path) -> str:
    try:
        import pdfplumber
    except ImportError as exc:  # pragma: no cover - environment issue
        raise UnreadableDocument("pdfplumber is not installed") from exc

    try:
        with pdfplumber.open(path) as pdf:
            pages = [page.extract_text() or "" for page in pdf.pages]
    except Exception as exc:
        raise UnreadableDocument(f"{path}: PDF could not be parsed ({exc})") from exc
    return "\n".join(pages)


_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _docx_text_of(element) -> str:
    """Text under one element (paragraph or cell), preserving line breaks.

    <w:br/> inside a run is a real line break: the xlsx+docx pairs put a
    party's name on the first line and its address on the following ones.
    Concatenating runs blindly welds them together ("...TRADINGON BEHALF
    OF..."), which both corrupts the name and defeats the first-line rule.
    """
    parts: list[str] = []
    for node in element.iter():
        if node.tag == f"{_W_NS}t":
            parts.append(node.text or "")
        elif node.tag in (f"{_W_NS}br", f"{_W_NS}cr"):
            parts.append("\n")
    lines = "".join(parts).split("\n")
    cleaned = [re.sub(r"[^\S\n]+", " ", line).strip() for line in lines]
    return "\n".join(line for line in cleaned if line)


def _read_docx(path: Path) -> str:
    """Read a .docx using only the standard library.

    A .docx is a zip whose word/document.xml holds the content, so this
    needs no third-party parser. That is deliberate: python-docx pulls in
    lxml, whose compiled _elementpath extension is blocked by Application
    Control policy on the dev machine — and a reader that only works on the
    deploy target is a reader that cannot be tested before deploying.
    """
    import zipfile
    from xml.etree import ElementTree

    try:
        with zipfile.ZipFile(path) as archive:
            xml = archive.read("word/document.xml")
    except (zipfile.BadZipFile, KeyError) as exc:
        raise UnreadableDocument(f"{path}: DOCX could not be parsed ({exc})") from exc

    try:
        root = ElementTree.fromstring(xml)
    except ElementTree.ParseError as exc:
        raise UnreadableDocument(f"{path}: DOCX XML is malformed ({exc})") from exc

    # Paragraphs living inside table cells are emitted as part of their row,
    # so they must not also be emitted standalone. ElementTree has no parent
    # links, so collect them up front.
    in_table = {
        id(paragraph)
        for row in root.iter(f"{_W_NS}tr")
        for paragraph in row.iter(f"{_W_NS}p")
    }

    lines: list[str] = []
    # Walk paragraphs and table rows in document order. Shipping docs put
    # most fields in tables: render each row as "label: value" so the same
    # alias matcher applies to tables and prose alike.
    for element in root.iter():
        if element.tag == f"{_W_NS}tr":
            cells = [
                text
                for cell in element.findall(f"{_W_NS}tc")
                if (text := _docx_text_of(cell))
            ]
            if len(cells) >= 2:
                # A label cell is single-line; a value cell may carry a
                # party name plus address lines, which must stay on their
                # own lines for the first-line rule to hold.
                label = cells[0].splitlines()[0]
                lines.append(f"{label}: {' '.join(cells[1:])}".split("\n", 1)[0])
                remainder = "\n".join(cells[1:]).split("\n", 1)
                if len(remainder) > 1:
                    lines.append(remainder[1])
            elif cells:
                lines.append(cells[0])
        elif element.tag == f"{_W_NS}p" and id(element) not in in_table:
            if text := _docx_text_of(element):
                lines.append(text)

    return "\n".join(lines)


def _read_xlsx(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment issue
        raise UnreadableDocument("openpyxl is not installed") from exc

    try:
        workbook = load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:
        raise UnreadableDocument(f"{path}: XLSX could not be parsed ({exc})") from exc

    lines: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
            if len(cells) >= 2:
                # These sheets pack a party's name and address into one cell
                # separated by "|", the same structure the docx pairs express
                # with <w:br/>. Normalizing both to real newlines is what lets
                # the first-line rule compare an xlsx SI against a docx BL.
                value = " ".join(cells[1:]).replace("|", "\n")
                value_lines = [part.strip() for part in value.split("\n") if part.strip()]
                lines.append(f"{cells[0]}: {value_lines[0] if value_lines else ''}")
                lines.extend(value_lines[1:])
            elif cells:
                lines.append(cells[0])
    workbook.close()
    return "\n".join(lines)


def extract_from_file(path: str | Path) -> dict[str, str]:
    """Convenience: read a document of any supported format and extract."""
    return extract_fields(read_document(path))
